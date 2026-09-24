"""THETIS-MRV workbook → report rows: the columns kept, read by position under a checked header."""

from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.mrv.enums import MrvSheet
from app.mrv.models import FIGURE_LABELS, MrvFile
from app.mrv.names import company_key, name_key

HEADER_ROW = 3
ETS_LAYOUT = {
    "imo": 0,
    "ship_name": 1,
    "ship_type": 2,
    "period_label": 3,
    "company_imo": 8,
    "company_name": 9,
    "co2_total": 28,
    "co2_between_ms": 29,
    "co2_departed_ms": 30,
    "co2_arrived_ms": 31,
    "co2_at_berth": 32,
    "co2_ets": 37,
}
"""From 2024: the company, the EU ETS figure, and full and partial reports on two sheets."""
PRE_ETS_LAYOUT = {
    "imo": 0,
    "ship_name": 1,
    "ship_type": 2,
    "period_label": 3,
    "co2_total": 24,
    "co2_between_ms": 25,
    "co2_departed_ms": 26,
    "co2_arrived_ms": 27,
    "co2_at_berth": 28,
}
"""2018 to 2023: one sheet of reports naming neither company nor EU ETS figure."""
LAYOUTS = (
    (ETS_LAYOUT, "co2_ets", "CO2 emissions to be reported under Directive 2003/87/EC [m tonnes]"),
    (PRE_ETS_LAYOUT, "co2_total", "Total CO₂ emissions [m tonnes]"),
)
"""Each layout, known by the header one of its columns carries."""


class MrvLayoutError(ValueError):
    """The workbook's columns are not where the parser reads them."""


def sheet_kind(title: str) -> MrvSheet:
    return MrvSheet.PARTIAL if "Partial" in title else MrvSheet.FULL


def sheet_layout(title: str, header: tuple[Any, ...]) -> dict[str, int]:
    for layout, column, expected in LAYOUTS:
        if header[layout[column]] == expected:
            return layout
    raise MrvLayoutError(f"{title}: its header matches no known layout")


def to_figure(value: Any) -> float | None:
    """A cell as tonnes, or None when EMSA left it blank or wrote text."""
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_workbook(content: bytes, file: MrvFile) -> list[dict[str, Any]]:
    """Every report in every sheet as a row for mrv_reports, what a layout lacks left empty."""
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    rows: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        header = next(sheet.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))
        layout = sheet_layout(sheet.title, header)
        for cells in sheet.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
            if not cells[layout["imo"]]:
                continue
            row = {name: cells[index] for name, index in layout.items()}
            company = row.get("company_name")
            rows.append(
                row
                | {name: to_figure(row.get(name)) for name in FIGURE_LABELS}
                | {
                    "imo": str(row["imo"]),
                    "period_label": str(row["period_label"]),
                    "company_imo": str(row["company_imo"]) if row.get("company_imo") else None,
                    "company_name": company,
                    "company_key": company_key(company) if company else None,
                    "ship_key": name_key(row["ship_name"]),
                    "sheet": sheet_kind(sheet.title),
                    "period": file.period,
                    "version": file.version,
                    "generated": file.generated,
                }
            )
    return rows
